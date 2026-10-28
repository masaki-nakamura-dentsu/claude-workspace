#!/usr/bin/env python3
"""外務省 海外安全情報オープンデータ → GBC出張管理用 危険レベル表(xlsx)

出力: GBCアレココ(GBCarekoko)サイトのドキュメント「事務局管理/GBC出張管理」の 外務省危険レベル.xlsx (Power Automate / Power BI が参照)
  T_国別 : 国名キー(別名含む) → 最大レベル / 既定レベル(都市が特定地域に該当しない場合) / 色 / URL
  T_地域 : 国名・レベル・対象地域の本文(参考)
  T_都市 : 国都市キー(国名|都市名) → 判定レベル  ※都市辞書.csv の行ごとに判定
都市辞書.csv (このスクリプトと同じフォルダ) に「国名,都市名,照合キーワード(;区切り)」を足すと都市単位の判定が増える。
会社Macでは社内プロキシ経由でのみ外務省サイトに届く。
"""
import csv, datetime, os, re, subprocess, sys, xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.expanduser(os.environ.get("MOFA_RISK_OUT", "~/MofaRisk/外務省危険レベル.xlsx"))  # ローカル控え(同期フォルダ外)
# アップロード先。OneDrive同期は反映されないことがあった(2026-10-20)ため、共有ブラウザ(CDP 9226)のログイン済みセッションでSharePoint RESTに直接上書きする
SP_SITE = "https://esq365.sharepoint.com/sites/GBCarekoko"
SP_FOLDER = "/sites/GBCarekoko/DocLib/事務局管理/GBC出張管理"
SP_NAME = "外務省危険レベル.xlsx"
SP_LIST = "7cc2eda0-a34a-4fdc-8057-3979cabd2628"  # GBC出張管理リスト
CDP = os.environ.get("MOFA_CDP", "http://127.0.0.1:9226")
PROXY = os.environ.get("MOFA_PROXY", "http://proxy1.dentsu.co.jp:8080")
BASE = "https://www.ezairyu.mofa.go.jp/opendata"
MOFA_TOP = "https://www.anzen.mofa.go.jp/"
COLORS = {0: "#FFFFFF", 1: "#FFFF00", 2: "#FFCC00", 3: "#FF0000", 4: "#8B0A03"}  # 外務省地図の凡例色
LABELS = {0: "レベルなし", 1: "レベル1:十分注意してください", 2: "レベル2:不要不急の渡航は止めてください",
          3: "レベル3:渡航は止めてください(渡航中止勧告)", 4: "レベル4:退避してください(退避勧告)"}
# Formsで入力されがちな表記 → 外務省の国名
ALIASES = {  # 外務省国名 → 追加の別名(「／」区切りの各部分と「（」より前は自動で別名になる)
    "アメリカ合衆国／米国（本土）": ["アメリカ", "USA", "US"], "英国／イギリス／グレートブリテンおよび北アイルランド連合王国": ["UK"],
    "アラブ首長国連邦": ["UAE"], "ロシア": ["ロシア連邦"],
}


EN = {}
_en = os.path.join(HERE, "英語国名.csv")
if os.path.exists(_en):
    for _r in csv.reader(open(_en, encoding="utf-8-sig")):
        if _r and _r[0].isdigit():
            EN[_r[0]] = [x.strip() for x in _r[1].split(";") if x.strip()]


def keys_for(name, cd=""):
    ks = [name] + ALIASES.get(name, []) + EN.get(cd, [])
    for part in name.split("／"):
        ks += [part, part.split("（")[0]]
    return list(dict.fromkeys(k.strip() for k in ks if k.strip()))
Z2H = str.maketrans("０１２３４５６７８９", "0123456789")

def get(url, stop=None, limit=5_000_000):
    """会社MacはPythonの証明書検証が社内プロキシで失敗するためcurlで取得。stop以降は読まずに打ち切る。"""
    cmd = ["curl", "-s", "-m", "120"] + (["-x", PROXY] if PROXY else []) + [url]
    p = subprocess.Popen(cmd, stdout=subprocess.PIPE)
    buf = b""
    try:
        while True:
            chunk = p.stdout.read(16384)
            if not chunk:
                break
            buf += chunk
            if stop and (stop in buf or len(buf) >= limit):
                break
    finally:
        if p.poll() is None:
            p.kill()
        p.wait()
    if not buf:
        raise RuntimeError(f"取得できません: {url}")
    return buf


def country_list():
    root = ET.fromstring(get(f"{BASE}/area/00A.xml"))
    return sorted({(c.findtext("cd"), c.findtext("name")) for c in root.iter("country") if c.findtext("cd")})


def head_of(cd):
    """国別XMLの先頭(個別メール本体の前)だけ読む。"""
    raw = get(f"{BASE}/country/{cd}A.xml", stop=b"<mail")
    txt = raw.decode("utf-8", "ignore")
    i = txt.find("<mail")
    txt = (txt[:i] if i > 0 else txt) + "</opendata>"
    return ET.fromstring(txt)


def parse_blocks(sub):
    """【危険レベル】節を (レベル, 地域) に分解。節の順序や「地域：レベル1」の1行書きにも対応。"""
    sub = (sub or "").translate(Z2H).replace("●", "\n●")
    m = re.search(r"【危険(レベル|度)】", sub)
    sect = sub[m.end():] if m else sub.split("【ポイント】")[0]
    j = sect.find("【")
    sect = sect[:j] if j >= 0 else sect
    blocks, head, last = [], "", ""
    for line in sect.splitlines():
        line = line.strip().replace("\u3000", "")
        if not line:
            continue
        if line.startswith("●"):
            head, last = line[1:], ""
            line = line[1:]
        m = re.search(r"レベル\s*([1-4])\s*(?=[：:《（(]|$)", line)
        if not m:
            last = line
            continue
        pre = line[:m.start()].strip("：: 　")
        region = pre or last or head
        if head and region != head:
            region = f"{head} {region}"
        blocks.append((int(m.group(1)), region))
        last = ""
    return blocks


def default_level(blocks):
    for lv, rg in blocks:
        if re.search(r"全土|全域$|上記以外|下記以外|上記を除く|その他の地域|それ以外の地域", rg):
            return lv
    return 0


def main():
    rows_c, rows_r, info, keymap = [], [], {}, {}
    for cd, name in country_list():
        name = name.strip()
        try:
            h = head_of(cd)
        except Exception as e:  # 1か国の失敗で全体を止めない
            print("skip", cd, name, e, file=sys.stderr)
            continue
        lv = max([n for n in (1, 2, 3, 4) if int(h.findtext(f"riskLevel{n}") or 0) > 0] or [0])
        blocks = parse_blocks(h.findtext("riskSubText"))
        dflt = default_level(blocks) if lv else 0
        # 国別の「危険・スポット・広域情報」ページ(地図詳細やryojiMailはブラウザで開くとエラーになるため使わない)。空だとSPのハイパーリンク列への書込みで失敗する
        ccd = (h.find("country").get("countryCd") if h.find("country") is not None else "") or ""
        url = f"https://www.anzen.mofa.go.jp/info/pcinfectionspothazardinfo_{int(ccd):03d}.html" if ccd.isdigit() else MOFA_TOP
        info[name] = (lv, dflt, blocks, url)
        for key in keys_for(name, cd):
            if key in keymap and "（本土）" not in name:
                continue  # 同じ別名は先勝ち(米国は本土を優先)
            keymap[key] = name
            rows_c = [r for r in rows_c if r[0] != key]
            rows_c.append([key, name, cd, lv, LABELS[lv], COLORS[lv], dflt, COLORS[dflt], url,
                           (h.findtext("riskTitle") or "")[:250]])
        for blv, rg in blocks:
            rows_r.append([name, blv, COLORS[blv], rg[:250]])

    # 日本(外務省データ対象外)は国内出張としてレベル0
    info["日本"] = (0, 0, [], MOFA_TOP)
    for key in ["日本", "Japan", "JAPAN"]:
        keymap[key] = "日本"
        rows_c.append([key, "日本", "0081", 0, "国内(対象外)", COLORS[0], 0, COLORS[0], MOFA_TOP, "国内出張"])

    # 都市辞書で都市単位の判定
    rows_t = []
    dic = os.path.join(HERE, "都市辞書.csv")
    for r in csv.DictReader(open(dic, encoding="utf-8-sig")) if os.path.exists(dic) else []:
        c = keymap.get(r["国名"].strip())
        if c not in info:
            print("都市辞書: 国名不明", r["国名"], file=sys.stderr)
            continue
        lv, dflt, blocks, url = info[c]
        keys = [r["都市名"].strip()] + [k.strip() for k in r.get("照合キーワード", "").split(";") if k.strip()]
        hits = [b for b in blocks if keys[0] in b[1]] or [b for b in blocks if any(k in b[1] for k in keys[1:])]
        if hits:
            clv, basis = max(b[0] for b in hits), "地域一致: " + max(hits)[1][:80]
        elif blocks or not lv:
            clv, basis = dflt, ("該当地域なし(国の既定)" if blocks else "国全体")
        else:
            clv, basis = lv, "地域を読み取れず(国内最大レベルを表示・要確認)"
        cities = [r["都市名"].strip()] + [x.strip() for x in (r.get("英語名") or "").split(";") if x.strip()]
        for ck in [k for k, v in keymap.items() if v == c]:
            for cn in cities:
                rows_t.append([f"{ck}|{cn}", c, r["都市名"].strip(), clv, LABELS[clv], COLORS[clv], basis, url])

    # SharePointリストの実際の申請(国名|都市名の生の文字列)ごとに判定。複数都市・県名付き等はトークンに分けて最大レベル
    done = {r[0] for r in rows_t}
    try:
        trips = list_trips()
    except Exception as e:
        print("リスト取得失敗(辞書のみで続行)", e, file=sys.stderr)
        trips = []
    for country, city in trips:
        key = f"{country}|{city}"
        c = keymap.get(country)
        if key in done or c not in info:
            continue
        done.add(key)
        lv, dflt, blocks, url = info[c]
        toks = [t for t in re.split(r"[、,，/／→・\s　及およびと&＆]+", city) if t]
        toks = [re.sub(r"(市|県|州|郡|区|町)$", "", t) for t in toks] + toks
        toks = [t for t in dict.fromkeys(toks) if len(t) >= 2]
        dic_hits = [r for r in rows_t if r[1] == c and r[2] in toks]
        hits = [b for b in blocks if any(t in b[1] for t in toks)]
        if c == "日本":
            clv, basis = 0, "国内出張"
        elif dic_hits:
            r = max(dic_hits, key=lambda x: x[3]); clv, basis = r[3], r[6]
        elif hits:
            clv, basis = max(b[0] for b in hits), "地域一致: " + max(hits)[1][:80]
        elif blocks or not lv:
            clv, basis = dflt, ("該当地域なし(国の既定)" if blocks else "国全体")
        else:
            clv, basis = lv, "地域を読み取れず(国内最大レベルを表示・要確認)"
        rows_t.append([key, c, city, clv, LABELS[clv], COLORS[clv], basis, url])

    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    wb = Workbook()
    def sheet(ws, name, header, rows):
        ws.title = name
        ws.append(header)
        for r in rows or [[""] * len(header)]:
            ws.append(r)
        ref = f"A1:{chr(64 + len(header))}{ws.max_row}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
    sheet(wb.active, "T_国別", ["CountryKey", "外務省国名", "国コード", "最大レベル", "レベル表示", "色", "既定レベル", "既定色", "危険情報URL", "見出し"], rows_c)
    sheet(wb.create_sheet(), "T_地域", ["外務省国名", "レベル", "色", "対象地域"], rows_r)
    sheet(wb.create_sheet(), "T_都市", ["CityKey", "外務省国名", "都市名", "レベル", "レベル表示", "色", "判定根拠", "危険情報URL"], rows_t)
    ws = wb.create_sheet("更新情報"); ws.append(["取得日時", stamp]); ws.append(["出典", "外務省 海外安全情報オープンデータ"])
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    upload(OUT)
    print(f"OK {stamp} 国キー{len(rows_c)} 地域{len(rows_r)} 都市{len(rows_t)} -> {OUT}")


def sp_eval(js, args):
    """共有ブラウザ(ログイン済み)のSharePointタブでJSを実行。"""
    import asyncio
    from playwright.async_api import async_playwright

    async def run():
        async with async_playwright() as p:
            b = await p.chromium.connect_over_cdp(CDP)
            ctx = b.contexts[0]
            page = await ctx.new_page()
            try:
                try:  # new_pageでウィンドウが前に出るので即最小化
                    s = await ctx.new_cdp_session(page)
                    w = await s.send("Browser.getWindowForTarget")
                    await s.send("Browser.setWindowBounds", {"windowId": w["windowId"], "bounds": {"windowState": "minimized"}})
                except Exception:
                    pass
                await page.goto(SP_SITE + "/_api/web?$select=Title", wait_until="domcontentloaded")
                return await page.evaluate(js, args)
            finally:
                await page.close()
    return asyncio.run(run())


def list_trips():
    """GBC出張管理リストの (国名, 都市名) を全件取得。"""
    js = """async ([site, list]) => {
      let u = site + "/_api/web/lists(guid'" + list + "')/items?$top=1000&$select=OData__x56fd__x540d_,OData__x90fd__x5e02__x540d_";
      const r = await fetch(u, {headers: {accept: "application/json;odata=nometadata"}});
      if (!r.ok) throw new Error("list " + r.status);
      return (await r.json()).value.map(x => [x.OData__x56fd__x540d_ || "", x.OData__x90fd__x5e02__x540d_ || ""]);
    }"""
    return [(a.strip(), b.strip()) for a, b in sp_eval(js, [SP_SITE, SP_LIST]) if a]


def upload(path):
    """SharePointの同じファイルに上書き(Files/add overwrite=true はファイルIDを保つので、Power AutomateのExcel接続が切れない)。"""
    import base64, time
    data = base64.b64encode(open(path, "rb").read()).decode()
    js = """async ([site, folder, name, b64]) => {
      const d = await fetch(site + "/_api/contextinfo", {method: "POST", headers: {accept: "application/json;odata=nometadata"}});
      if (!d.ok) return "digest " + d.status;
      const digest = (await d.json()).FormDigestValue;
      const bin = Uint8Array.from(atob(b64), c => c.charCodeAt(0));
      const u = site + "/_api/web/GetFolderByServerRelativePath(decodedurl='" + folder + "')/Files/AddUsingPath(decodedurl='" + name + "',overwrite=true)";
      const r = await fetch(u, {method: "POST", headers: {"X-RequestDigest": digest, accept: "application/json;odata=nometadata"}, body: bin});
      return r.status + " " + (await r.text()).slice(0, 300);
    }"""
    last = ""
    for i in range(8):  # Excel Online等で開かれてロック中(423)なら待って再試行(最大約16分)
        last = sp_eval(js, [SP_SITE, SP_FOLDER, SP_NAME, data])
        if last.startswith("200"):
            print("uploaded", last[:120])
            return
        time.sleep(120)
    raise RuntimeError("SharePointアップロード失敗: " + last)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        sys.path.insert(0, os.path.join(HERE, "..", "notify"))
        try:
            from teams_notify import notify
            notify("ERROR", "外務省危険Lv", "取得失敗", detail=repr(e))
        except Exception:
            pass
        raise
